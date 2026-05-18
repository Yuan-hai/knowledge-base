"""Document parser — extract text from PDF and DOCX files."""
from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file using PyPDF2."""
    from PyPDF2 import PdfReader

    logger.info("Parsing PDF: %s", file_path)
    reader = PdfReader(file_path)
    text_parts: list[str] = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
        else:
            logger.warning("No text extracted from page %d of %s", i + 1, file_path)
    full_text = "\n".join(text_parts)
    if not full_text.strip():
        raise ValueError(f"No extractable text found in PDF: {file_path}")
    logger.info("PDF parsed: %d pages, %d characters", len(reader.pages), len(full_text))
    return full_text


def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document

    logger.info("Parsing DOCX: %s", file_path)
    doc = Document(file_path)
    text_parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    full_text = "\n".join(text_parts)
    if not full_text.strip():
        raise ValueError(f"No extractable text found in DOCX: {file_path}")
    logger.info("DOCX parsed: %d paragraphs, %d characters", len(doc.paragraphs), len(full_text))
    return full_text


def parse_document(file_path: str) -> str:
    """Auto-detect document type by extension and extract text.

    Args:
        file_path: Path to the document file.

    Returns:
        Extracted text content.

    Raises:
        ValueError: If the file extension is unsupported or text extraction fails.
    """
    ext = Path(file_path).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported document format '{ext}'. Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    # Should never reach here due to check above, but safety fallback
    raise ValueError(f"Unsupported format: {ext}")
